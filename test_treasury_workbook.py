#!/usr/bin/env python3
"""
Treasury Calendar Workbook Test Suite
Validates structure, formulas, and data integrity
"""

import openpyxl
from openpyxl import load_workbook
from datetime import datetime
import sys

def test_workbook():
    """Run comprehensive tests on the Treasury Calendar Workbook"""

    print("=" * 70)
    print("TREASURY CALENDAR WORKBOOK - TEST SUITE")
    print("=" * 70)

    filepath = '/home/user/Wedxdev/Treasury_Calendar_Workbook.xlsx'

    try:
        wb = load_workbook(filepath, data_only=False)
        print(f"\n✓ Workbook loaded successfully: {filepath}")
    except Exception as e:
        print(f"\n✗ FAILED to load workbook: {e}")
        return False

    tests_passed = 0
    tests_failed = 0

    # =========================================================================
    # TEST 1: Verify all expected sheets exist
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 1: Sheet Structure Validation")
    print("-" * 50)

    expected_sheets = [
        'Config', 'FX_Rates', 'SRC_Collections', 'SRC_OPeX', 'SRC_CAPEX',
        'SRC_HR', 'SRC_Debt', 'SRC_Fixed', 'SRC_Variable', 'SRC_Deposits',
        'Vendor_Master', 'Payment_Data', 'Treasury_Calendar', 'Payment_Calendar'
    ]

    for sheet_name in expected_sheets:
        if sheet_name in wb.sheetnames:
            print(f"  ✓ Sheet '{sheet_name}' exists")
            tests_passed += 1
        else:
            print(f"  ✗ Sheet '{sheet_name}' MISSING")
            tests_failed += 1

    print(f"\nActual sheets: {len(wb.sheetnames)}")
    print(f"Expected sheets: {len(expected_sheets)}")

    # =========================================================================
    # TEST 2: Config Sheet Validation
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 2: Config Sheet Validation")
    print("-" * 50)

    ws = wb['Config']

    # Check title
    if ws['A1'].value == 'TREASURY CALENDAR - CONFIGURATION':
        print("  ✓ Config title correct")
        tests_passed += 1
    else:
        print(f"  ✗ Config title incorrect: {ws['A1'].value}")
        tests_failed += 1

    # Check base currency
    if ws['B4'].value == 'SAR':
        print("  ✓ Base currency is SAR")
        tests_passed += 1
    else:
        print(f"  ✗ Base currency incorrect: {ws['B4'].value}")
        tests_failed += 1

    # Check weekend configuration
    if ws['B9'].value == 'Friday' and ws['B10'].value == 'Saturday':
        print("  ✓ Weekend days configured (Friday/Saturday)")
        tests_passed += 1
    else:
        print(f"  ✗ Weekend days incorrect")
        tests_failed += 1

    # =========================================================================
    # TEST 3: FX Rates Sheet Validation
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 3: FX Rates Sheet Validation")
    print("-" * 50)

    ws = wb['FX_Rates']

    # Check currencies exist
    currencies = ['SAR', 'USD', 'AED', 'EUR', 'GBP', 'SGD']
    found_currencies = []
    for row in range(6, 12):
        ccy = ws.cell(row=row, column=1).value
        if ccy in currencies:
            found_currencies.append(ccy)

    if len(found_currencies) == 6:
        print(f"  ✓ All 6 currencies found: {found_currencies}")
        tests_passed += 1
    else:
        print(f"  ✗ Missing currencies. Found: {found_currencies}")
        tests_failed += 1

    # Check SAR rate is 1.0
    sar_rate = ws.cell(row=6, column=2).value
    if sar_rate == 1.0:
        print("  ✓ SAR rate is 1.0000")
        tests_passed += 1
    else:
        print(f"  ✗ SAR rate incorrect: {sar_rate}")
        tests_failed += 1

    # =========================================================================
    # TEST 4: Payment_Data Sheet Validation (Single Source of Truth)
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 4: Payment_Data Sheet Validation")
    print("-" * 50)

    ws = wb['Payment_Data']

    # Check title
    if 'PAYMENT DATA' in str(ws['A1'].value):
        print("  ✓ Payment_Data title correct")
        tests_passed += 1
    else:
        print(f"  ✗ Payment_Data title incorrect")
        tests_failed += 1

    # Check for Single Source of Truth notice
    found_ssot = False
    for row in range(1, 10):
        cell_value = str(ws.cell(row=row, column=1).value or '')
        if 'SINGLE SOURCE OF TRUTH' in cell_value or 'MASTER' in cell_value:
            found_ssot = True
            break

    if found_ssot:
        print("  ✓ Single Source of Truth notice present")
        tests_passed += 1
    else:
        print("  ✗ Single Source of Truth notice missing")
        tests_failed += 1

    # Check for sample payment data
    payment_count = 0
    for row in range(6, 50):
        payment_id = ws.cell(row=row, column=1).value
        if payment_id and str(payment_id).startswith('PAY-'):
            payment_count += 1

    if payment_count >= 20:
        print(f"  ✓ Sample payments found: {payment_count}")
        tests_passed += 1
    else:
        print(f"  ✗ Insufficient sample payments: {payment_count}")
        tests_failed += 1

    # Check for IFERROR formulas in SAR amount column
    formula_count = 0
    for row in range(6, 35):
        cell = ws.cell(row=row, column=9)
        if cell.value and str(cell.value).startswith('=IFERROR'):
            formula_count += 1

    if formula_count >= 10:
        print(f"  ✓ IFERROR formulas present: {formula_count}")
        tests_passed += 1
    else:
        print(f"  ✗ Insufficient IFERROR formulas: {formula_count}")
        tests_failed += 1

    # =========================================================================
    # TEST 5: Vendor_Master Sheet Validation
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 5: Vendor_Master Sheet Validation")
    print("-" * 50)

    ws = wb['Vendor_Master']

    # Check for vendor categories
    categories_found = set()
    for row in range(4, 60):
        category = ws.cell(row=row, column=3).value
        if category:
            categories_found.add(category)

    expected_categories = {'VENDOR', 'HR', 'GOVT', 'STRATEGIC', 'LOAN', 'OPEX'}
    if expected_categories.issubset(categories_found):
        print(f"  ✓ All vendor categories found: {expected_categories}")
        tests_passed += 1
    else:
        print(f"  ✗ Missing categories. Found: {categories_found}")
        tests_failed += 1

    # Check for vendor IDs
    vendor_ids = []
    for row in range(4, 60):
        vid = ws.cell(row=row, column=1).value
        if vid and not str(vid).startswith('Vendor'):
            vendor_ids.append(vid)

    if len(vendor_ids) >= 25:
        print(f"  ✓ Sufficient vendors: {len(vendor_ids)}")
        tests_passed += 1
    else:
        print(f"  ✗ Insufficient vendors: {len(vendor_ids)}")
        tests_failed += 1

    # =========================================================================
    # TEST 6: Treasury_Calendar Sheet Validation
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 6: Treasury_Calendar Sheet Validation")
    print("-" * 50)

    ws = wb['Treasury_Calendar']

    # Check title
    if ws['A1'].value == 'TREASURY CALENDAR':
        print("  ✓ Treasury Calendar title correct")
        tests_passed += 1
    else:
        print(f"  ✗ Treasury Calendar title incorrect: {ws['A1'].value}")
        tests_failed += 1

    # Check date header rows exist
    date_labels = ['YEAR', 'MONTH', 'WEEK #', 'DAY', 'DATE', 'WEEKDAY']
    labels_found = 0
    for row in range(2, 8):
        label = ws.cell(row=row, column=1).value
        if label in date_labels:
            labels_found += 1

    if labels_found >= 5:
        print(f"  ✓ Date header rows present: {labels_found}/6")
        tests_passed += 1
    else:
        print(f"  ✗ Missing date headers: {labels_found}/6")
        tests_failed += 1

    # Check for financial categories
    categories = ['OPENING BALANCE', 'INFLOWS', 'OUTFLOWS', 'NET CASH FLOW', 'CLOSING BALANCE']
    found_cats = []
    for row in range(8, 40):
        cat = ws.cell(row=row, column=1).value
        if cat and cat in categories:
            found_cats.append(cat)

    if len(found_cats) >= 4:
        print(f"  ✓ Financial categories found: {found_cats}")
        tests_passed += 1
    else:
        print(f"  ✗ Missing financial categories: {found_cats}")
        tests_failed += 1

    # Check for SUMIF formulas (data aggregation from source sheets)
    sumif_count = 0
    for row in range(10, 30):
        for col in range(4, 20):
            cell = ws.cell(row=row, column=col)
            if cell.value and 'SUMIF' in str(cell.value):
                sumif_count += 1
                break
        if sumif_count >= 5:
            break

    if sumif_count >= 3:
        print(f"  ✓ SUMIF formulas for source aggregation present")
        tests_passed += 1
    else:
        print(f"  ✗ Insufficient SUMIF formulas")
        tests_failed += 1

    # Check horizontal timeline (60 days)
    date_columns = 0
    for col in range(4, 70):
        cell = ws.cell(row=6, column=col)
        if cell.value and isinstance(cell.value, datetime):
            date_columns += 1

    if date_columns >= 50:
        print(f"  ✓ Horizontal timeline: {date_columns} date columns")
        tests_passed += 1
    else:
        print(f"  ✗ Insufficient date columns: {date_columns}")
        tests_failed += 1

    # =========================================================================
    # TEST 7: Payment_Calendar Sheet Validation
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 7: Payment_Calendar Sheet Validation")
    print("-" * 50)

    ws = wb['Payment_Calendar']

    # Check title
    if ws['A1'].value == 'PAYMENT CALENDAR':
        print("  ✓ Payment Calendar title correct")
        tests_passed += 1
    else:
        print(f"  ✗ Payment Calendar title incorrect: {ws['A1'].value}")
        tests_failed += 1

    # Check for vendor category headers
    cat_headers = ['VENDOR PAYMENTS', 'HUMAN RESOURCES', 'GOVERNMENT PAYMENTS',
                   'STRATEGIC PAYMENTS', 'LOAN REPAYMENTS', 'OPERATING EXPENSES']
    found_headers = []
    for row in range(9, 50):
        val = ws.cell(row=row, column=1).value
        if val in cat_headers:
            found_headers.append(val)

    if len(found_headers) >= 5:
        print(f"  ✓ Payment category headers found: {len(found_headers)}")
        tests_passed += 1
    else:
        print(f"  ✗ Missing category headers: {found_headers}")
        tests_failed += 1

    # Check for SUMPRODUCT formulas (payment lookup)
    sumproduct_count = 0
    for row in range(10, 45):
        for col in range(5, 15):
            cell = ws.cell(row=row, column=col)
            if cell.value and 'SUMPRODUCT' in str(cell.value):
                sumproduct_count += 1
                break
        if sumproduct_count >= 5:
            break

    if sumproduct_count >= 3:
        print(f"  ✓ SUMPRODUCT formulas for payment lookup present")
        tests_passed += 1
    else:
        print(f"  ✗ Insufficient SUMPRODUCT formulas")
        tests_failed += 1

    # Check for conditional formatting (status-based highlighting)
    cf_rules = len(ws.conditional_formatting._cf_rules)
    if cf_rules >= 3:
        print(f"  ✓ Conditional formatting rules: {cf_rules}")
        tests_passed += 1
    else:
        print(f"  ✗ Insufficient conditional formatting: {cf_rules}")
        tests_failed += 1

    # =========================================================================
    # TEST 8: Source Sheet Formulas Validation
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 8: Source Sheet Formulas Validation")
    print("-" * 50)

    source_sheets = ['SRC_Collections', 'SRC_OPeX', 'SRC_CAPEX', 'SRC_HR',
                     'SRC_Debt', 'SRC_Fixed', 'SRC_Variable', 'SRC_Deposits']

    for sheet_name in source_sheets:
        ws = wb[sheet_name]

        # Check for SAR conversion formula with IFERROR
        has_formula = False
        for row in range(4, 15):
            for col in range(5, 10):
                cell = ws.cell(row=row, column=col)
                if cell.value and 'IFERROR' in str(cell.value) and 'VLOOKUP' in str(cell.value):
                    has_formula = True
                    break
            if has_formula:
                break

        if has_formula:
            print(f"  ✓ {sheet_name}: IFERROR+VLOOKUP formula present")
            tests_passed += 1
        else:
            print(f"  ✗ {sheet_name}: Missing conversion formula")
            tests_failed += 1

    # =========================================================================
    # TEST 9: Data Validation Rules
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 9: Data Validation Rules")
    print("-" * 50)

    ws = wb['Payment_Data']
    dv_count = len(ws.data_validations.dataValidation)

    if dv_count >= 2:
        print(f"  ✓ Data validation rules in Payment_Data: {dv_count}")
        tests_passed += 1
    else:
        print(f"  ✗ Insufficient data validation rules: {dv_count}")
        tests_failed += 1

    # =========================================================================
    # TEST 10: Weekend Highlighting (Saudi Calendar)
    # =========================================================================
    print("\n" + "-" * 50)
    print("TEST 10: Weekend Highlighting (Saudi Calendar)")
    print("-" * 50)

    ws = wb['Treasury_Calendar']

    # Check for Friday/Saturday in weekday row
    weekday_row = 7
    friday_found = False
    saturday_found = False

    for col in range(4, 70):
        day = ws.cell(row=weekday_row, column=col).value
        if day == 'Fri':
            friday_found = True
        elif day == 'Sat':
            saturday_found = True

    if friday_found and saturday_found:
        print("  ✓ Friday and Saturday weekdays present in timeline")
        tests_passed += 1
    else:
        print(f"  ✗ Weekend days not found (Fri: {friday_found}, Sat: {saturday_found})")
        tests_failed += 1

    # =========================================================================
    # SUMMARY
    # =========================================================================
    print("\n" + "=" * 70)
    print("TEST SUMMARY")
    print("=" * 70)
    print(f"\n  Tests Passed: {tests_passed}")
    print(f"  Tests Failed: {tests_failed}")
    print(f"  Total Tests:  {tests_passed + tests_failed}")
    print(f"\n  Success Rate: {tests_passed / (tests_passed + tests_failed) * 100:.1f}%")

    if tests_failed == 0:
        print("\n  ✓ ALL TESTS PASSED - Workbook is valid!")
        return True
    else:
        print(f"\n  ✗ {tests_failed} TESTS FAILED - Please review")
        return False


if __name__ == '__main__':
    success = test_workbook()
    sys.exit(0 if success else 1)
